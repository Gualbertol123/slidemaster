"""PowerPoint checks for the .pptx export (docs/next/07 §4.3, 05 PX-01…PX-03; S11.1). Dev only, like pdfcheck.py: the
structure check is the standard library; the text, table and chart checks need python-pptx 1.0.2 (MIT); the open
check needs PowerPoint (Windows, COM through PowerShell). The helper stays standard-library only.

    python tools/pptxcheck.py check a.pptx [b.pptx ...] [--slides N] [--no-ppt] [--json out.json] [--sources sources.json] [--expect expect.json]
        (1) structure: the zip opens; [Content_Types].xml types every part; every relationship's target exists and none
            points outside the file (no links); every XML part parses; the slide count (presentation.xml, = N when
            given); the slide size 12 192 000 × 6 858 000 EMU; no empty a:t; no connector shapes (p:cxnSp); every
            slide has a layout, every layout the master; the pictures are PNG or JPEG; and a lint of what PowerPoint
            repairs without saying so in automation (S11 review: an unknown or misplaced child, a duplicate shape id,
            a negative size, a short table row pass Presentations.Open): the schema's child order of sp, pic, spPr, ln,
            effectLst, blipFill, txBody, p, pPr, rPr, tc, tcPr; shape ids, placeholders and relationship ids unique per
            part; value ranges (ext, rot, angles, gradient stops, font sizes, spacing, insets, margins, shadows,
            preset adjustments, custom paths); every table row as many cells as the grid has columns
        (2) content (python-pptx): each slide's texts (text frames, table cells, chart titles), tables (has_table),
            charts (has_chart) with their cached categories and values, pictures
        (3) PowerPoint (unless --no-ppt, or no PowerPoint): Presentations.Open read-only without a window – a file
            it cannot read raises in automation (not every repair does: (1)'s lint covers those) –; the slide count and every shape's text read back and
            compared with (2)
        (4) with --sources ({file: {workbook, charts: [{slide, sheet, range, start?, before?, name?, axis?, mask?}]}}): one series of the
            chart (named `name`, when given) holds the workbook's cells at their category positions – the cells' first at
            category `start` (default 0) plus the visible cells of the ranges `before` ([{sheet, range}]: the category
            groups' before its own), a blank cell a gap there, and no value outside them
            (openpyxl, the values Excel saved; hidden cells skipped as the chart skips them). A pie may instead hold the
            cells' numbers with blanks and text left out (the old reading)
        (5) real objects, gated (S14.P; AUDIT-2 VER-3, PX-02 / PX-03, 09 §7.2): with --expect ({file: [{slide, tables,
            charts}]}, tools/real-objects.mjs: what the core's own export draws) each slide holds exactly those real
            tables (python-pptx has_table) and native charts (has_chart), a slide not listed none;
            where PowerPoint is present, it must open the file (3) and see the same tables (HasTable) and charts
            (HasChart) per slide as python-pptx – a failure, never a count. Each report says `ppt`: opened / failed /
            absent (no PowerPoint here: 1, 2, 4 and python-pptx's part of 5 run) / skipped (--no-ppt). Also in (1): no
            chart formula (c:f) names another file (09 §4.2: no external links)
        exit 1 when any check fails; --json writes every file's report (the content of (2) included)
    python tools/pptxcheck.py info a.pptx [...]      the report of (1) and (2) as JSON, nothing checked
"""
import argparse
import json
import os
import posixpath
import re
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
}
SLIDE_SIZE = (12192000, 6858000)
RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"


A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]
# the schema's child order (ECMA-376 part 1) of the elements the writer makes; "r", "br", "fld" are one slot of a:p
ORDER = {
    "spPr": ["xfrm", "custGeom", "prstGeom", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill", "ln", "effectLst", "effectDag", "scene3d", "sp3d", "extLst"],
    "rPr": ["ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill", "effectLst", "effectDag", "highlight", "uLnTx", "uLn", "uFillTx", "uFill", "latin", "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst"],
    "tcPr": ["lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr", "cell3D", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill", "headers", "extLst"],
    "pPr": ["lnSpc", "spcBef", "spcAft", "buClrTx", "buClr", "buSzTx", "buSzPct", "buSzPts", "buFontTx", "buFont", "buNone", "buAutoNum", "buChar", "buBlip", "tabLst", "defRPr", "extLst"],
    "effectLst": ["blur", "fillOverlay", "glow", "innerShdw", "outerShdw", "prstShdw", "reflection", "softEdge"],
    "ln": ["noFill", "solidFill", "gradFill", "pattFill", "prstDash", "custDash", "round", "bevel", "miter", "headEnd", "tailEnd", "extLst"],
    "blipFill": ["blip", "srcRect", "tile", "stretch"],
    "sp": ["nvSpPr", "spPr", "style", "txBody", "extLst"],
    "pic": ["nvPicPr", "blipFill", "spPr", "style", "extLst"],
    "tc": ["txBody", "tcPr", "extLst"],
    "txBody": ["bodyPr", "lstStyle", "p"],
    "p": ["pPr", "r", "r", "r", "endParaRPr"],
}
SLOT = {"br": "r", "fld": "r"}
C = "{%s}" % NS["c"]
# the chart part's (c:) child order of the elements the writer makes
C_ORDER = {
    "pieChart": ["varyColors", "ser", "dLbls", "firstSliceAng", "extLst"],
    "doughnutChart": ["varyColors", "ser", "dLbls", "firstSliceAng", "holeSize", "extLst"],
    "dPt": ["idx", "invertIfNegative", "marker", "bubble3D", "explosion", "spPr", "pictureOptions", "extLst"],
    "dLbls": ["dLbl", "delete", "numFmt", "spPr", "txPr", "dLblPos", "showLegendKey", "showVal", "showCatName", "showSerName", "showPercent", "showBubbleSize", "separator", "showLeaderLines", "leaderLines", "extLst"],
    "dLbl": ["idx", "delete", "layout", "tx", "numFmt", "spPr", "txPr", "dLblPos", "showLegendKey", "showVal", "showCatName", "showSerName", "showPercent", "showBubbleSize", "separator", "extLst"],
    "chart": ["title", "autoTitleDeleted", "pivotFmts", "view3D", "floor", "sideWall", "backWall", "plotArea", "legend", "plotVisOnly", "dispBlanksAs", "showDLblsOverMax", "extLst"],
    "chartSpace": ["date1904", "lang", "roundedCorners", "AlternateContent", "clrMapOvr", "pivotSource", "protection", "chart", "spPr", "txPr", "externalData", "printSettings", "userShapes", "extLst"],
    "legend": ["legendPos", "legendEntry", "layout", "overlay", "spPr", "txPr", "extLst"],
    # column and bar charts (4.7, S14.2; 09 §4.2): the bar group, its axes, a manual layout, a title, a hidden legend
    # entry (the plot area and each chart group's c:ser: chart_order below)
    "barChart": ["barDir", "grouping", "varyColors", "ser", "dLbls", "gapWidth", "overlap", "serLines", "axId", "extLst"],
    "catAx": ["axId", "scaling", "delete", "axPos", "majorGridlines", "minorGridlines", "title", "numFmt", "majorTickMark", "minorTickMark", "tickLblPos", "spPr", "txPr",
              "crossAx", "crosses", "crossesAt", "auto", "lblAlgn", "lblOffset", "tickLblSkip", "tickMarkSkip", "noMultiLvlLbl", "extLst"],
    "valAx": ["axId", "scaling", "delete", "axPos", "majorGridlines", "minorGridlines", "title", "numFmt", "majorTickMark", "minorTickMark", "tickLblPos", "spPr", "txPr",
              "crossAx", "crosses", "crossesAt", "crossBetween", "majorUnit", "minorUnit", "dispUnits", "extLst"],
    "scaling": ["logBase", "orientation", "max", "min", "extLst"],
    "manualLayout": ["layoutTarget", "xMode", "yMode", "wMode", "hMode", "x", "y", "w", "h", "extLst"],
    "title": ["tx", "layout", "overlay", "spPr", "txPr", "extLst"],
    "legendEntry": ["idx", "delete", "txPr", "extLst"],
    # line and area charts, a line's marker, multi-level categories (4.7, S14.3; 09 §4.2)
    "lineChart": ["grouping", "varyColors", "ser", "dLbls", "dropLines", "hiLowLines", "upDownBars", "marker", "smooth", "axId", "extLst"],
    "areaChart": ["grouping", "varyColors", "ser", "dLbls", "dropLines", "axId", "extLst"],
    "marker": ["symbol", "size", "spPr", "extLst"],
    "multiLvlStrRef": ["f", "multiLvlStrCache", "extLst"],
    "multiLvlStrCache": ["ptCount", "lvl", "extLst"],
    # scatter charts (4.7, S14.4; 09 §4.2)
    "scatterChart": ["scatterStyle", "varyColors", "ser", "dLbls", "axId", "extLst"],
}
LN_TAGS = ("ln", "lnL", "lnR", "lnT", "lnB")


def _local(tag):
    return tag.split("}")[1] if "}" in tag else tag


def lint(name, root):
    """the problems of one slide, layout or master part that PowerPoint would repair (see the module's doc)"""
    out = []

    def bad(msg):
        if len(out) < 20:
            out.append("%s: %s" % (name, msg))
    ids = [e.get("id") for e in root.iter(P + "cNvPr")]
    if len(ids) != len(set(ids)):
        bad("duplicate shape ids")
    phs = [(e.get("type"), e.get("idx")) for e in root.iter(P + "ph")]
    if len(phs) != len(set(phs)):
        bad("duplicate placeholders")

    def num(e, k):
        try:
            return int(e.get(k))
        except (TypeError, ValueError):
            bad("<%s %s=%r> is not a number" % (_local(e.tag), k, e.get(k)))
            return None
    for e in root.iter():
        t, ns = _local(e.tag), e.tag[:e.tag.index("}") + 1] if "}" in e.tag else ""
        key = "ln" if t in LN_TAGS else t
        if key in ORDER and ns in (A, P) and not (t in ("sp", "pic") and ns != P):
            seq, last = ORDER[key], -1
            for ch in e:
                n = SLOT.get(_local(ch.tag), _local(ch.tag))
                if n not in seq:
                    bad("<%s> in <%s>" % (n, t))
                    continue
                i = seq.index(n)
                if i < last:
                    bad("<%s> after a later child in <%s>" % (n, t))
                last = max(last, i)
        if ns == C and t in C_ORDER:
            seq, last = C_ORDER[t], -1
            for ch in e:
                n = _local(ch.tag)
                if n not in seq:
                    bad("<c:%s> in <c:%s>" % (n, t))
                    continue
                if seq.index(n) < last:
                    bad("<c:%s> after a later child in <c:%s>" % (n, t))
                last = max(last, seq.index(n))
        if ns != A:
            continue
        if t == "off":
            num(e, "x"), num(e, "y")
        elif t == "ext" and e.get("cx") is not None:
            cx, cy = num(e, "cx"), num(e, "cy")
            if cx is not None and cy is not None and (cx < 0 or cy < 0):
                bad("negative size %s × %s" % (cx, cy))
        elif t == "xfrm" and e.get("rot") is not None:
            v = num(e, "rot")
            if v is not None and not 0 <= v < 21600000:
                bad("rotation %s" % v)
        elif t == "lin":
            v = num(e, "ang")
            if v is not None and not 0 <= v < 21600000:
                bad("gradient angle %s" % v)
        elif t == "gs":
            v = num(e, "pos")
            if v is not None and not 0 <= v <= 100000:
                bad("gradient stop at %s" % v)
        elif t == "gsLst" and len(list(e)) < 2:
            bad("a gradient with fewer than two stops")
        elif t == "path" and e.get("w") is not None:
            w, h = num(e, "w"), num(e, "h")
            if w is not None and h is not None and (w <= 0 or h <= 0):
                bad("a custom path of size %s × %s" % (w, h))
        elif t == "gd":
            m = re.match(r"val (-?\d+)$", e.get("fmla") or "")
            if not m:
                bad("adjustment %r" % e.get("fmla"))
            elif e.get("name") == "adj" and not 0 <= int(m.group(1)) <= 50000:
                bad("adj %s" % m.group(1))
        elif t in ("rPr", "endParaRPr", "defRPr") and e.get("sz") is not None:
            v = num(e, "sz")
            if v is not None and not 100 <= v <= 400000:
                bad("font size %s" % v)
        elif t == "spcPts":
            v = num(e, "val")
            if v is not None and not 0 <= v <= 158400:
                bad("spacing %s" % v)
        elif t in ("outerShdw", "innerShdw"):
            vals = [num(e, k) for k in ("blurRad", "dist", "dir")]
            if None not in vals and (vals[0] < 0 or vals[1] < 0 or not 0 <= vals[2] < 21600000):
                bad("shadow %s" % vals)
        elif t == "bodyPr":
            if any(e.get(k) is not None and (num(e, k) or 0) < 0 for k in ("lIns", "tIns", "rIns", "bIns")):
                bad("a negative inset")
        elif t == "tcPr":
            if any(e.get(k) is not None and (num(e, k) or 0) < 0 for k in ("marL", "marR", "marT", "marB")):
                bad("a negative cell margin")
        elif t == "tbl":
            cols = len(list(e.iter(A + "gridCol")))
            for tr in e.findall(A + "tr"):
                if len(tr.findall(A + "tc")) != cols:
                    bad("a table row of %d cells in a grid of %d" % (len(tr.findall(A + "tc")), cols))
                    break
    return out


def _rels_of(part):
    d, f = posixpath.split(part)
    return posixpath.join(d, "_rels", f + ".rels")


def structure(path, slides=None):
    """the structure check (1): a list of problems (empty = passed) and facts"""
    bad, facts = [], {"bytes": os.path.getsize(path)}
    try:
        z = zipfile.ZipFile(path)
    except zipfile.BadZipFile as e:
        return ["not a zip: %s" % e], facts
    names = set(z.namelist())
    if "[Content_Types].xml" not in names:
        return ["no [Content_Types].xml"], facts
    ct = ET.fromstring(z.read("[Content_Types].xml"))
    defaults = {d.get("Extension").lower(): d.get("ContentType") for d in ct.findall("ct:Default", NS)}
    overrides = {o.get("PartName").lstrip("/"): o.get("ContentType") for o in ct.findall("ct:Override", NS)}
    for o in overrides:
        if o not in names:
            bad.append("content type for a missing part: " + o)
    xml = {}
    for n in sorted(names):
        if n.endswith("/") or n == "[Content_Types].xml":
            continue
        ext = n.rsplit(".", 1)[-1].lower()
        if n not in overrides and ext not in defaults:
            bad.append("no content type: " + n)
        if ext in ("xml", "rels"):
            try:
                xml[n] = ET.fromstring(z.read(n))
            except ET.ParseError as e:
                bad.append("XML does not parse: %s (%s)" % (n, e))
    # relationships: targets exist, nothing external, ids unique
    for n, root in xml.items():
        if re.match(r"ppt/(slides|slideLayouts|slideMasters|charts)/[^/]+\.xml$", n):
            bad.extend(lint(n, root))
            bad.extend(chart_order(n, root))            # (1) M14: per parent (its own function below)
        if not n.endswith(".rels"):
            continue
        rids = [r.get("Id") for r in root.findall("rel:Relationship", NS)]
        if len(rids) != len(set(rids)):
            bad.append("duplicate relationship ids: " + n)
        base = posixpath.dirname(posixpath.dirname(n))
        for r in root.findall("rel:Relationship", NS):
            if r.get("TargetMode") == "External":
                bad.append("a link outside the file: %s → %s" % (n, r.get("Target")))
                continue
            t = posixpath.normpath(posixpath.join(base, r.get("Target"))).lstrip("/")
            if t not in names:
                bad.append("missing target: %s → %s" % (n, r.get("Target")))
    pres = xml.get("ppt/presentation.xml")
    if pres is None:
        return bad + ["no ppt/presentation.xml"], facts
    sz = pres.find("p:sldSz", NS)
    if sz is None or (int(sz.get("cx")), int(sz.get("cy"))) != SLIDE_SIZE:
        bad.append("slide size %s, not 16:9 12192000 × 6858000" % ((None if sz is None else (sz.get("cx"), sz.get("cy"))),))
    ids = pres.findall("p:sldIdLst/p:sldId", NS)
    facts["slides"] = len(ids)
    if slides is not None and len(ids) != slides:
        bad.append("%d slides, expected %d" % (len(ids), slides))
    prels = {r.get("Id"): r for r in xml.get("ppt/_rels/presentation.xml.rels", ET.Element("x")).findall("rel:Relationship", NS)}
    order = []
    for s in ids:
        r = prels.get(s.get("{%s}id" % NS["r"]))
        if r is None:
            bad.append("a slide id without its relationship")
            continue
        order.append(posixpath.normpath(posixpath.join("ppt", r.get("Target"))))
    facts["masters"] = len([n for n in names if re.match(r"ppt/slideMasters/[^/]+\.xml$", n)])
    facts["layouts"] = len([n for n in names if re.match(r"ppt/slideLayouts/[^/]+\.xml$", n)])
    if facts["masters"] < 1:
        bad.append("no slide master")
    for part in [n for n in names if re.match(r"ppt/slideLayouts/[^/]+\.xml$", n)]:
        rel = xml.get(_rels_of(part))
        if rel is None or not any(r.get("Type") == RT + "slideMaster" for r in rel.findall("rel:Relationship", NS)):
            bad.append("a layout without its master: " + part)
    media = [n for n in names if n.startswith("ppt/media/")]
    for m in media:
        head = z.read(m)[:8]
        if not (head.startswith(b"\x89PNG") or head.startswith(b"\xff\xd8")):
            bad.append("a picture that is not PNG or JPEG: " + m)
    facts["media"] = len(media)
    facts["media_bytes"] = sum(z.getinfo(m).file_size for m in media)
    for i, part in enumerate(order):
        root = xml.get(part)
        if root is None:
            bad.append("slide %d: no part %s" % (i + 1, part))
            continue
        rel = xml.get(_rels_of(part))
        if rel is None or not any(r.get("Type") == RT + "slideLayout" for r in rel.findall("rel:Relationship", NS)):
            bad.append("slide %d: no layout" % (i + 1))
        if root.find(".//p:cxnSp", NS) is not None:
            bad.append("slide %d: a connector shape" % (i + 1))
        # (S14.P, 09 §4.2: no external links) a chart's formula never names another file – every c:f points into its own
        # embedded workbook
        for ch in [posixpath.normpath(posixpath.join(posixpath.dirname(part), r.get("Target"))) for r in (rel.findall("rel:Relationship", NS) if rel is not None else [])
                   if r.get("Type") == RT + "chart"]:
            croot = xml.get(ch)
            for f in (croot.iter("{%s}f" % NS["c"]) if croot is not None else []):
                # another workbook is named as [n] / [Book.xlsx] or by a path; Excel forbids [ ] / and the backslash in
                # sheet names, so a sheet called like a file ('Q3.xlsx'!A1) does not trip it (S14.P review, 10)
                if re.search(r"[\[\]\\/]", f.text or ""):
                    bad.append("slide %d: a chart formula names a file: %s" % (i + 1, f.text))
                    break
        for t in root.iter("{%s}t" % NS["a"]):
            if not (t.text or ""):
                bad.append("slide %d: an empty a:t" % (i + 1))
                break
    facts["parts"] = len(names)
    return bad, facts


def _shape_texts(shape, out, info):
    """texts of a shape (groups walked), tables' cells and charts' titles; counts of tables, charts, pictures"""
    if shape.shape_type == 6 or getattr(shape, "shapes", None) is not None and shape.__class__.__name__ == "GroupShape":
        for s in shape.shapes:
            _shape_texts(s, out, info)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        info["tables"] += 1
        cells = []
        for row in shape.table.rows:
            for cell in row.cells:
                if cell.is_spanned:
                    continue
                t = cell.text_frame.text
                cells.append(t)
                if t.strip():
                    out.append(t)
        info["cells"].append(cells)
        return
    if getattr(shape, "has_chart", False) and shape.has_chart:
        ch = shape.chart
        plot = ch.plots[0]
        info["charts"].append({"type": str(ch.chart_type), "title": ch.chart_title.text_frame.text if ch.has_title else "",
                               "categories": [str(c) for c in plot.categories], "values": list(plot.series[0].values),
                               # (4.7, S14.2) every series of every plot: its name and cached values (None = a gap)
                               "series": [dict({"name": s.name, "values": list(s.values)}, **scatter_xs(s)) for pl in ch.plots for s in pl.series],
                               # (4.7, S14.3) each chart group's kind (BarPlot, LinePlot, AreaPlot …) and the category levels
                               "plots": [type(pl).__name__ for pl in ch.plots], "levels": plot.categories.depth})
        return
    if shape.shape_type == 13:
        info["pictures"] += 1
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
        out.append(shape.text_frame.text)
    if getattr(shape, "is_placeholder", False) and shape.is_placeholder:
        info["placeholders"].append(str(shape.placeholder_format.type).split(".")[-1].split(" ")[0])


def content(path):
    """the content (2): per slide its texts, tables' cells, charts, pictures, placeholders, layout"""
    from pptx import Presentation
    prs = Presentation(path)
    slides = []
    for s in prs.slides:
        texts, info = [], {"tables": 0, "cells": [], "charts": [], "pictures": 0, "placeholders": []}
        for sh in s.shapes:
            _shape_texts(sh, texts, info)
        slides.append({"layout": s.slide_layout.name, "texts": texts, **info})
    return {"slides": slides, "layouts": [l.name for l in prs.slide_layouts], "masters": len(prs.slide_masters)}


PS_OPEN = r"""
$ErrorActionPreference = "Stop"
$files = Get-Content -Raw -Encoding UTF8 $args[0] | ConvertFrom-Json
# a PowerPoint that was running before (someone's, another check's) is used and left running
$was = $true; try { [Runtime.InteropServices.Marshal]::GetActiveObject("PowerPoint.Application") | Out-Null } catch { $was = $false }
$app = New-Object -ComObject PowerPoint.Application
$out = @()
function texts($shapes, [ref]$acc, [ref]$nt, [ref]$nc) {
  foreach ($sh in $shapes) {
    if ($sh.Type -eq 6) { texts $sh.GroupItems $acc $nt $nc; continue }
    if ($sh.HasChart) { $nc.Value++ }
    if ($sh.HasTable) { $nt.Value++; $t = $sh.Table; for ($r = 1; $r -le $t.Rows.Count; $r++) { for ($c = 1; $c -le $t.Columns.Count; $c++) { $x = $t.Cell($r, $c).Shape.TextFrame.TextRange.Text; if ($x.Trim()) { $acc.Value += $x } } }; continue }
    if ($sh.HasTextFrame -and $sh.TextFrame.HasText) { $acc.Value += $sh.TextFrame.TextRange.Text }
  }
}
try {
  foreach ($f in $files) {
    $r = @{ file = $f; ok = $false; slides = 0; texts = @(); tables = @(); charts = @(); error = "" }
    try {
      $p = $app.Presentations.Open($f, -1, 0, 0)
      $r.slides = $p.Slides.Count
      foreach ($s in $p.Slides) { $acc = @(); $nt = 0; $nc = 0; texts $s.Shapes ([ref]$acc) ([ref]$nt) ([ref]$nc); $r.texts += ,@($acc); $r.tables += $nt; $r.charts += $nc }
      $p.Close(); $r.ok = $true
    } catch { $r.error = $_.Exception.Message }
    $out += $r
  }
} finally { if (-not $was) { $app.Quit() | Out-Null } }
ConvertTo-Json -Depth 6 -InputObject @($out) | Out-File -Encoding UTF8 $args[1]
"""


def powerpoint_available():
    if os.name != "nt":
        return False
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-Command", "[bool]([Type]::GetTypeFromProgID('PowerPoint.Application'))"],
                           capture_output=True, text=True, timeout=60)
        return r.stdout.strip() == "True"
    except (OSError, subprocess.TimeoutExpired):
        return False


def powerpoint_open(paths, timeout=600):
    """(3): each file opened in PowerPoint (read-only, no window): opened or the error, the slide count, the texts.
    One PowerShell call per file, each within `timeout` s (reading back a large deck's table cells takes a while over
    COM): a file that does not come back in time is a failure of its own, the others are still checked"""
    d = tempfile.mkdtemp(prefix="pptxcheck-")
    script = os.path.join(d, "open.ps1")
    with open(script, "w", encoding="utf-8-sig") as f:
        f.write(PS_OPEN)
    out = []
    for k, p in enumerate(paths):
        lst, res = os.path.join(d, "files%d.json" % k), os.path.join(d, "out%d.json" % k)
        with open(lst, "w", encoding="utf-8") as f:
            json.dump([os.path.abspath(p)], f)
        try:
            subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script, lst, res], capture_output=True, text=True, timeout=timeout)
            with open(res, encoding="utf-8-sig") as f:
                r = json.load(f)
            out.extend(r if isinstance(r, list) else [r])
        except subprocess.TimeoutExpired:
            out.append({"file": os.path.abspath(p), "ok": False, "slides": 0, "texts": [], "error": "not read back within %d s" % timeout})
        except (OSError, ValueError) as e:
            out.append({"file": os.path.abspath(p), "ok": False, "slides": 0, "texts": [], "error": "no answer: %s" % e})
    return out


def squash(s):
    import unicodedata
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", s or ""))


def source_values(workbook, sheet, rng):
    """the numbers of a range as a chart reads them: rows then columns, hidden rows / columns, text and blanks skipped"""
    import openpyxl
    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb[sheet]
    out = []
    cells = ws[rng]
    # (a single cell is a Cell, a row or a column a tuple of cells, a rectangle a tuple of rows)
    for row in (cells if isinstance(cells, tuple) else ((cells,),)):
        for c in (row if isinstance(row, tuple) else (row,)):
            if ws.row_dimensions[c.row].hidden or ws.column_dimensions[c.column_letter].hidden:
                continue
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                out.append(float(c.value))
    return out


def source_cells(workbook, sheet, rng):
    """a range's cells as a chart reads them, by position: rows then columns, hidden rows / columns skipped, a number as a
    float, anything else (blank, text) None"""
    import openpyxl
    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb[sheet]
    cells = ws[rng]
    out = []
    for row in (cells if isinstance(cells, tuple) else ((cells,),)):
        for c in (row if isinstance(row, tuple) else (row,)):
            if ws.row_dimensions[c.row].hidden or ws.column_dimensions[c.column_letter].hidden:
                continue
            out.append(float(c.value) if isinstance(c.value, (int, float)) and not isinstance(c.value, bool) else None)
    return out


def scatter_xs(series):
    """(4.7, S14.4) a scatter series' cached x values (None = no point), read from its c:xVal – python-pptx reads only the
    y values; {} for any other series"""
    xv = series._element.find(C + "xVal")
    if xv is None:
        return {}
    cache = xv.find(".//" + C + "numCache")
    if cache is None:
        return {"xs": []}
    n = int(cache.find(C + "ptCount").get("val"))
    out = [None] * n
    for pt in cache.findall(C + "pt"):
        out[int(pt.get("idx"))] = float(pt.find(C + "v").text)
    return {"xs": out}


def series_at(values, cells, start):
    """(4) a series' cached values hold the cells at categories start… (a blank cell is a gap) and nothing elsewhere (4.7,
    S14.3 review 4: a series cached under the wrong categories – off by its group's start – fails)"""
    # (cells past the cache's end may only be gaps: a cache's ptCount can stop at its last point)
    if start < 0 or len(values) < start + len(cells) and any(w is not None for w in cells[max(0, len(values) - start):]):
        return False
    for i, v in enumerate(values):
        w = cells[i - start] if start <= i < start + len(cells) else None
        if (v is None) != (w is None):
            return False
        if w is not None and abs(float(v) - w) > 1e-12 * max(1.0, abs(w)):
            return False
    return True


# ---- M14 (4.7; S14.2 review, 4): the chart part's order checks that depend on the parent, its own function ----
# a choice of the schema: one rank for every chart group and every axis of a plot area
C_SLOT = dict([(k, "CHART") for k in ("areaChart", "area3DChart", "lineChart", "line3DChart", "stockChart", "radarChart", "scatterChart", "pieChart", "pie3DChart",
                                      "doughnutChart", "barChart", "bar3DChart", "ofPieChart", "surfaceChart", "surface3DChart", "bubbleChart")]
              + [(k, "AXIS") for k in ("valAx", "catAx", "dateAx", "serAx")])
# each chart group's own series type (ECMA-376 part 1, CT_PieSer, CT_BarSer, CT_LineSer, CT_AreaSer): a pie's c:ser takes
# no invertIfNegative, a bar's no explosion, an area's no marker
_PIE_SER = ["idx", "order", "tx", "spPr", "explosion", "dPt", "dLbls", "cat", "val", "extLst"]
C_PARENT_ORDER = {
    "plotArea": ["layout", "CHART", "AXIS", "dTable", "spPr", "extLst"],
    "pieChart/ser": _PIE_SER,
    "doughnutChart/ser": _PIE_SER,
    "barChart/ser": ["idx", "order", "tx", "spPr", "invertIfNegative", "pictureOptions", "dPt", "dLbls", "trendline", "errBars", "cat", "val", "shape", "extLst"],
    "lineChart/ser": ["idx", "order", "tx", "spPr", "marker", "dPt", "dLbls", "trendline", "errBars", "cat", "val", "smooth", "extLst"],
    "areaChart/ser": ["idx", "order", "tx", "spPr", "pictureOptions", "dPt", "dLbls", "trendline", "errBars", "cat", "val", "extLst"],
    # (4.7, S14.4) CT_ScatterSer
    "scatterChart/ser": ["idx", "order", "tx", "spPr", "marker", "dPt", "dLbls", "trendline", "errBars", "xVal", "yVal", "smooth", "extLst"],
}


def chart_order(name, root):
    """(1) a chart part's children whose order depends on the parent (4.7, S14.2 review 4): the plot area's (its chart
    groups, then its axes) and each chart group's c:ser, by the group's own series type; a c:ser under any other group is
    not checked. The problems as lint() words them"""
    out = []
    if root.tag != C + "chartSpace":
        return out

    def check(e, key, label):
        seq, last = C_PARENT_ORDER[key], -1
        for ch in e:
            n = _local(ch.tag)
            n = C_SLOT.get(n, n) if key == "plotArea" else n
            if n not in seq:
                out.append("%s: <c:%s> in <c:%s>" % (name, n, label))
                continue
            if seq.index(n) < last:
                out.append("%s: <c:%s> after a later child in <c:%s>" % (name, n, label))
            last = max(last, seq.index(n))
    for pa in root.iter(C + "plotArea"):
        check(pa, "plotArea", "plotArea")
        for g in pa:
            key = _local(g.tag) + "/ser"
            if key in C_PARENT_ORDER:
                for ser in g.findall("c:ser", NS):
                    check(ser, key, "ser> of <c:" + _local(g.tag))
    return out[:20]


PS_PLOTS = r"""
$ErrorActionPreference = "Stop"
$job = Get-Content -Raw -Encoding UTF8 $args[0] | ConvertFrom-Json
$was = $true; try { [Runtime.InteropServices.Marshal]::GetActiveObject("PowerPoint.Application") | Out-Null } catch { $was = $false }
$app = New-Object -ComObject PowerPoint.Application
$out = @{ file = $job.file; ok = $false; charts = @(); error = "" }
function charts($shapes, $n, [ref]$acc) {
  # (S14.6 review 5) the arrow and line marks beside a chart in its group: their boxes and flips (their ends)
  $lines = @()
  foreach ($sh in $shapes) {
    if ($sh.Type -ne 6 -and $sh.Name -match '^(Arrow|Line) \d+$') { $lines += @{ name = $sh.Name; left = $sh.Left; top = $sh.Top; width = $sh.Width; height = $sh.Height; flipH = [int]$sh.HorizontalFlip; flipV = [int]$sh.VerticalFlip } }
  }
  foreach ($sh in $shapes) {
    if ($sh.Type -eq 6) { charts $sh.GroupItems $n $acc; continue }
    if (-not $sh.HasChart) { continue }
    $ch = $sh.Chart; $pa = $ch.PlotArea
    $c = @{ slide = $n; left = $sh.Left; top = $sh.Top; width = $sh.Width; height = $sh.Height; inLeft = $pa.InsideLeft; inTop = $pa.InsideTop; inWidth = $pa.InsideWidth; inHeight = $pa.InsideHeight; ymin = $null; ymax = $null; y2min = $null; y2max = $null; lines = $lines }
    try { $ax = $ch.Axes(2, 1); $c.ymin = $ax.MinimumScale; $c.ymax = $ax.MaximumScale } catch {}
    try { $ax = $ch.Axes(2, 2); $c.y2min = $ax.MinimumScale; $c.y2max = $ax.MaximumScale } catch {}
    $acc.Value += $c
  }
}
try {
  $p = $app.Presentations.Open($job.file, -1, 0, 0)
  $acc = @()
  foreach ($s in $p.Slides) {
    charts $s.Shapes $s.SlideIndex ([ref]$acc)
    if ($job.png) { $s.Export((Join-Path $job.png ("slide-" + $s.SlideIndex + ".png")), "PNG", 1600, 900) }
  }
  $out.charts = $acc; $p.Close(); $out.ok = $true
} catch { $out.error = $_.Exception.Message } finally { if (-not $was) { $app.Quit() | Out-Null } }
ConvertTo-Json -Depth 6 -InputObject $out | Out-File -Encoding UTF8 $args[1]
"""


def plots(path, expect, png=None, timeout=600):
    """PX-06 (4.7, S14.6; 09 §4.2), Windows with PowerPoint only: the marks over a native chart land on its data when
    PowerPoint's plot is the slide's plot – each chart of `expect` ([{slide, frame, plot, val, val2?, arrows}], slide px;
    written by core/tests/pptx-charts-m14.test.ts beside each file as <name>.plots.json) is found by its frame, and
    PowerPoint's inner plot area (PlotArea.Inside*, points; 1 px = 0.6 pt) must be within 2 px of `plot`, its value axes'
    bounds (MinimumScale / MaximumScale) the drawn ones, and each arrow / line mark's two ends (its line shape's box and
    flips) within 2 px of the slide's `a` and `b` (compare_plots). `png`: a folder for every slide exported by PowerPoint at 1600 × 900
    (Slide.Export) to look at. Returns (problems, PowerPoint's charts)"""
    if not powerpoint_available():
        return ["PowerPoint is not available here (PX-06 runs on Windows with PowerPoint)"], []
    d = tempfile.mkdtemp(prefix="pptxcheck-plots-")
    script, job, res = os.path.join(d, "plots.ps1"), os.path.join(d, "job.json"), os.path.join(d, "out.json")
    with open(script, "w", encoding="utf-8-sig") as f:
        f.write(PS_PLOTS)
    if png:
        os.makedirs(png, exist_ok=True)
    with open(job, "w", encoding="utf-8") as f:
        json.dump({"file": os.path.abspath(path), "png": os.path.abspath(png) if png else ""}, f)
    try:
        subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script, job, res], capture_output=True, text=True, timeout=timeout)
        with open(res, encoding="utf-8-sig") as f:
            r = json.load(f)
    except (OSError, ValueError, subprocess.TimeoutExpired) as e:
        return ["PowerPoint: no answer (%s)" % e], []
    if not r.get("ok"):
        return ["PowerPoint: %s" % (r.get("error") or "not opened")], []
    got = r.get("charts") or []
    got = got if isinstance(got, list) else [got]
    return compare_plots(expect, got), got


def _px(pt):
    return (pt or 0) / 0.6


def line_ends(ln):
    """a line shape's two ends in slide px from PowerPoint's Left / Top / Width / Height (points) and its flips: unflipped
    it runs from the box's top-left to its bottom-right; HorizontalFlip / VerticalFlip (msoTrue = -1) swap the sides –
    as pptx/chart-overlay.ts writes it (flipH when b is left of a, flipV when b is above a)"""
    x0, y0 = _px(ln.get("left")), _px(ln.get("top"))
    x1, y1 = x0 + _px(ln.get("width")), y0 + _px(ln.get("height"))
    if ln.get("flipH"):
        x0, x1 = x1, x0
    if ln.get("flipV"):
        y0, y1 = y1, y0
    return (x0, y0), (x1, y1)


def compare_marks(arrows, lines, tol=2.0):
    """(S14.6 review 5) each expected arrow / line mark ({a, b}, slide px) against PowerPoint's line shapes of the chart's
    group (line_ends): one shape each, its first end on `a` and its second on `b` within `tol` px (the heads sit on the
    ends, so a swapped line is a wrong arrow); no shape left over. Returns the problems"""
    lines = lines if isinstance(lines, list) else ([lines] if lines else [])
    left, bad = list(range(len(lines))), []
    ends = [line_ends(l) for l in lines]
    for k, e in enumerate(arrows or []):
        a, b = (e["a"]["x"], e["a"]["y"]), (e["b"]["x"], e["b"]["y"])
        off = lambda j: max(abs(ends[j][0][0] - a[0]), abs(ends[j][0][1] - a[1]), abs(ends[j][1][0] - b[0]), abs(ends[j][1][1] - b[1]))
        best = min(left, key=off) if left else None
        if best is None or off(best) > tol:
            near = "" if best is None else ": the nearest, %s, runs %s → %s (%.1f px off)" % (lines[best].get("name", "?"), tuple(round(v, 1) for v in ends[best][0]), tuple(round(v, 1) for v in ends[best][1]), off(best))
            bad.append("mark %d %s → %s: no line shape on it%s" % (k + 1, a, b, near))
            continue
        left.remove(best)
    for j in left:
        bad.append("%s: a line shape no mark of the slide has" % lines[j].get("name", "?"))
    return bad


def compare_plots(expect, got):
    """PX-06's comparison (plots(), separate so that it runs without PowerPoint): `expect` as <name>.plots.json, `got`
    PowerPoint's charts as PS_PLOTS reports them (points). Returns the problems"""
    px = _px
    bad = []
    for k, e in enumerate(expect):
        F = e["frame"]
        mine = [c for c in got if c["slide"] == e["slide"] and abs(px(c["left"]) - F["x"]) <= 2 and abs(px(c["top"]) - F["y"]) <= 2]
        if not mine:
            bad.append("chart %d (slide %d): no native chart at its frame (%.0f, %.0f)" % (k + 1, e["slide"], F["x"], F["y"]))
            continue
        c, P = mine[0], e["plot"]
        at = {"x": px(c["left"] + c["inLeft"]), "y": px(c["top"] + c["inTop"]), "w": px(c["inWidth"]), "h": px(c["inHeight"])}
        off = max(abs(at[q] - P[q]) for q in "xywh")
        if off > 2:
            bad.append("chart %d (slide %d): PowerPoint's plot %s is %.1f px off the slide's %s" % (k + 1, e["slide"], {q: round(v, 1) for q, v in at.items()}, off, P))
        for key, lo, hi in (("val", "ymin", "ymax"), ("val2", "y2min", "y2max")):
            V = e.get(key)
            if V and (c.get(lo) is None or abs(c[lo] - V["min"]) > 1e-9 * max(1, abs(V["min"])) or abs(c[hi] - V["max"]) > 1e-9 * max(1, abs(V["max"]))):
                bad.append("chart %d (slide %d): PowerPoint's %s axis %s–%s, drawn %s–%s" % (k + 1, e["slide"], "second" if key == "val2" else "value", c.get(lo), c.get(hi), V["min"], V["max"]))
        # (S14.6 review 5: the marks – the plan's "within 2 px at 1600" – each arrow's ends on the slide's)
        bad.extend("chart %d (slide %d): %s" % (k + 1, e["slide"], m) for m in compare_marks(e.get("arrows") or [], c.get("lines") or []))
    return bad


def real_objects(rep_content, expect, who):
    """(5) the gate of AUDIT-2 VER-3 / PX-03 (S14.P, 09 §7.2): each slide has exactly the real tables (has_table) and
    native charts (has_chart) the core's own export draws there (`expect`, tools/real-objects.mjs: a table of a sheet the
    workbook lacks, a hidden slide are not drawn, so not expected) – never a picture of one, never one too many; a slide
    `expect` does not list must hold none. `who` names the reader in the message"""
    bad = []
    slides = rep_content["slides"] if rep_content else []
    want = {e["slide"]: e for e in expect or []}
    for n in sorted(set(want) | set(range(1, len(slides) + 1))):
        s = slides[n - 1] if 0 < n <= len(slides) else None
        for k, word in (("tables", "real table"), ("charts", "native chart")):
            w = (want.get(n) or {}).get(k) or 0
            have = (len(s["charts"]) if k == "charts" else s["tables"]) if s else 0
            if have != w:
                bad.append("slide %d: %s %d %s(s), the export draws %d" % (n, who, have, word, w))
        # (4.7, S14.6, PX-05: tools/real-charts.mjs) each native chart's chart groups, as the core writes them
        kinds = (want.get(n) or {}).get("kinds")
        if kinds is not None and s is not None:
            got = sorted(chart_kinds(c) for c in s["charts"])
            if got != sorted(kinds):
                bad.append("slide %d: %s chart types %s, the export writes %s" % (n, who, got, sorted(kinds)))
    return bad


PLOT_KIND = {"BarPlot": "bar", "LinePlot": "line", "AreaPlot": "area", "XyPlot": "scatter", "PiePlot": "pie", "DoughnutPlot": "doughnut"}


def chart_kinds(chart):
    """a chart of the content (2) as tools/real-charts.mjs names it: its plots' kinds in order, joined ("bar+line"). (5)
    compares them per slide, sorted, when an --expect entry has `kinds` (4.7, S14.6, PX-05: every chart type native)"""
    return "+".join(PLOT_KIND.get(p, p) for p in chart.get("plots") or [])


def check(paths, slides=None, ppt=True, report=None, sources=None, expect=None):
    """`expect` ({file: [{slide, tables, charts}]}): the real tables and native charts each slide must hold (5), exactly. Where
    PowerPoint is present it is a gate, not a count (AUDIT-2 VER-3, 09 §7.2): the file must open, and PowerPoint must see
    as many tables (HasTable) and charts (HasChart) per slide as python-pptx does; without it (CI, Linux) the structural
    checks and python-pptx's counts run, and each report says `ppt: absent`"""
    ok, reports = True, []
    use_ppt = ppt and powerpoint_available()
    opened = {os.path.abspath(r["file"]).lower(): r for r in powerpoint_open(paths)} if use_ppt else {}
    for p in paths:
        bad, facts = structure(p, slides)
        rep = {"file": p, **facts, "structure": bad or "ok", "ppt": "opened" if use_ppt else "absent" if ppt else "skipped"}
        try:
            rep["content"] = content(p)
        except Exception as e:                      # noqa: BLE001 (a file python-pptx cannot read is a failure)
            bad = bad + ["python-pptx cannot read it: %s" % e]
            rep["content"] = None
        exp = (expect or {}).get(p, (expect or {}).get(os.path.abspath(p)))
        if exp is not None:
            bad = bad + real_objects(rep["content"], exp, "python-pptx reads")
        src = (sources or {}).get(p) or (sources or {}).get(os.path.abspath(p))
        if src and rep["content"]:
            for k, spec in enumerate(src["charts"]):
                found = [c for s in rep["content"]["slides"][spec["slide"] - 1:spec["slide"]] for c in s["charts"]]
                want = source_values(src["workbook"], spec["sheet"], spec["range"])
                cells = source_cells(src["workbook"], spec["sheet"], spec["range"])
                # (4.7, S14.4 review 5: `mask` {sheet, range} – a scatter series' x cells: a y cell whose x is blank or text
                # has no point, so no cache)
                if spec.get("mask"):
                    m = source_cells(src["workbook"], spec["mask"]["sheet"], spec["mask"]["range"])
                    cells = [c if k < len(m) and m[k] is not None else None for k, c in enumerate(cells)]
                # (a group's start: `start` plus the visible cells of the category ranges `before` it – hidden rows and
                # columns skipped, as the chart skips them; S14.3 re-check)
                start = int(spec.get("start") or 0) + sum(len(source_cells(src["workbook"], b["sheet"], b["range"])) for b in spec.get("before") or [])
                name = spec.get("name")
                # (the cache keeps Excel's 15 significant digits: equal within that)
                same = lambda a, b: len(a) == len(b) and all(abs(float(x) - y) <= 1e-12 * max(1.0, abs(y)) for x, y in zip(a, b))
                # (4.7, S14.3: one series of the chart – a combo's line, a group's columns – at its categories, from its
                # group's start; a pie as before)
                pie = lambda c: "PIE" in c["type"] or "DOUGHNUT" in c["type"]
                # (4.7, S14.4: `axis: "x"` – a scatter series' x values, against its own c:xVal cache)
                key = "xs" if spec.get("axis") == "x" else "values"
                mine = [c for c in found
                        if any(key in x and series_at(list(x[key]), cells, start) and (name is None or x.get("name") == name) for x in (c.get("series") or [{"values": c["values"]}]))
                        or (pie(c) and not start and name is None and same(c["values"], want))]
                if not mine:
                    bad.append("chart %d (slide %d, %s!%s%s): no series with the cells' values %s at categories %d…"
                               % (k + 1, spec["slide"], spec["sheet"], spec["range"], " '%s'" % name if name else "", cells[:4], start))
        if use_ppt:
            o = opened.get(os.path.abspath(p).lower())
            if not o or not o.get("ok"):
                rep["ppt"] = "failed"
                bad.append("PowerPoint: %s" % ((o or {}).get("error") or "not opened"))
            else:
                as_list = lambda v: v if isinstance(v, list) else ([] if v is None else [v])     # (ConvertTo-Json: one slide is a number)
                rep["powerpoint"] = {"slides": o["slides"], "tables": as_list(o.get("tables")), "charts": as_list(o.get("charts"))}
                if o["slides"] != facts.get("slides"):
                    bad.append("PowerPoint reads %s slides, the file has %s" % (o["slides"], facts.get("slides")))
                if rep["content"]:
                    # (5) PowerPoint sees every real table and native chart python-pptx sees (VER-3: gated, not counted)
                    pt, pc = rep["powerpoint"]["tables"], rep["powerpoint"]["charts"]
                    for i, s in enumerate(rep["content"]["slides"]):
                        if (pt[i] if i < len(pt) else 0) != s["tables"] or (pc[i] if i < len(pc) else 0) != len(s["charts"]):
                            bad.append("slide %d: PowerPoint sees %s table(s) and %s chart(s), python-pptx %d and %d"
                                       % (i + 1, pt[i] if i < len(pt) else 0, pc[i] if i < len(pc) else 0, s["tables"], len(s["charts"])))
                    for i, (mine, theirs) in enumerate(zip(rep["content"]["slides"], o["texts"] or [])):
                        # (as sets: PowerPoint gives a merged cell's text for every cell it covers)
                        a, b = set(squash(t) for t in mine["texts"]), set(squash(t) for t in (theirs or []))
                        if a != b:
                            bad.append("slide %d: PowerPoint reads other texts (%d vs %d)" % (i + 1, len(b), len(a)))
        rep["ok"] = not bad
        rep["problems"] = bad
        reports.append(rep)
        ok = ok and not bad
        print(("PASS " if not bad else "FAIL ") + os.path.basename(p) + ("" if not bad else ": " + "; ".join(bad[:4])))
    if not use_ppt and ppt:
        print("(PowerPoint is not available here: checks 1–2, 4 and python-pptx's real tables and charts (5) only; on a PC with PowerPoint, 3 and 5 are gates)")
    if report:
        with open(report, "w", encoding="utf-8") as f:
            json.dump(reports, f, ensure_ascii=False, indent=1)
    return ok


def main(argv):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("check")
    c.add_argument("files", nargs="+")
    c.add_argument("--slides", type=int)
    c.add_argument("--no-ppt", action="store_true")
    c.add_argument("--json")
    c.add_argument("--sources")
    c.add_argument("--expect")
    i = sub.add_parser("info")
    i.add_argument("files", nargs="+")
    m = sub.add_parser("plots", help="PX-06 (S14.6): PowerPoint's plot areas and axis bounds against the slide's (Windows)")
    m.add_argument("files", nargs="+")
    m.add_argument("--png", help="a folder: every slide exported by PowerPoint (1600 × 900)")
    a = ap.parse_args(argv)
    if a.cmd == "plots":
        ok = True
        for p in a.files:
            with open(re.sub(r"\.pptx$", "", p) + ".plots.json", encoding="utf-8") as f:
                bad, _ = plots(p, json.load(f), os.path.join(a.png, os.path.splitext(os.path.basename(p))[0]) if a.png else None)
            print(("PASS " if not bad else "FAIL ") + os.path.basename(p) + ("" if not bad else ": " + "; ".join(bad[:4])))
            ok = ok and not bad
        return 0 if ok else 1
    if a.cmd == "info":
        out = []
        for p in a.files:
            bad, facts = structure(p)
            out.append({"file": p, **facts, "structure": bad or "ok", "content": content(p)})
        json.dump(out, sys.stdout, ensure_ascii=False, indent=1)
        return 0
    sources = None
    if a.sources:
        with open(a.sources, encoding="utf-8") as f:
            sources = json.load(f)
    expect = None
    if a.expect:
        with open(a.expect, encoding="utf-8") as f:
            expect = json.load(f)
    return 0 if check(a.files, a.slides, not a.no_ppt, a.json, sources, expect) else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
