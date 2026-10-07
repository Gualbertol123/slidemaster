"""Generates the test workbooks (python3 make_fixtures.py). Needs openpyxl (dev machines only)."""
import datetime, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule

HERE = os.path.dirname(os.path.abspath(__file__))
thin, med = Side(style="thin", color="000000"), Side(style="medium", color="1F3864")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
navy = PatternFill("solid", fgColor="1F3864"); grey = PatternFill("solid", fgColor="D9D9D9"); pale = PatternFill("solid", fgColor="DDEBF7")
green = PatternFill("solid", fgColor="C6EFCE"); red = PatternFill("solid", fgColor="FFC7CE")

def report():
    wb = Workbook(); ws = wb.active; ws.title = "SLIDE_1"
    ws["A1"] = "IBD – Total Banks Loans & Deposits"; ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "Weekly update, figures in EUR m"
    ws["B3"] = "x"; ws["K20"] = "x"
    hdr = ["Bank", "Loans", "Δ w/w", "Deposits", "Δ w/w", "Share", "Since", "Note"]
    for i, h in enumerate(hdr):
        c = ws.cell(row=4, column=3 + i, value=h); c.fill = navy; c.font = Font(bold=True, color="FFFFFF"); c.alignment = Alignment(horizontal="center"); c.border = box
    ws.merge_cells("C5:J5"); ws["C5"] = "Domestic"; ws["C5"].fill = grey; ws["C5"].font = Font(bold=True)
    banks = ["Alpha", "Beta", "Gamma", "Delta", "Epsilon", "Zeta", "Eta", "Theta", "Iota", "Kappa", "Lambda", "Mu", "Nu"]
    for k, b in enumerate(banks):
        r = 6 + k
        vals = [b, 1234567.891 * (k + 1), (-1) ** k * 2.5 * (k + 1), 987654.3 * (k + 2), (-1) ** (k + 1) * 1.25 * k, 0.0123 * (k + 1), datetime.date(2024, 1 + k % 12, 1 + k), "ok" if k % 3 else "check"]
        fmts = [None, "#,##0", "+0.0;-0.0;0.0", "#,##0.0,,\" m\"", "0.00", "0.0%", "dd/mm/yyyy", None]
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=3 + i, value=v); c.border = box
            if fmts[i]: c.number_format = fmts[i]
    ws.row_dimensions[9].hidden = True                      # hidden row inside the table
    ws.column_dimensions["H"].hidden = True                 # hidden column
    r = 19; ws.cell(row=r, column=3, value="TOTAL").font = Font(bold=True)
    for i in range(1, 8):
        c = ws.cell(row=r, column=3 + i); c.fill = pale; c.font = Font(bold=True); c.border = Border(top=med, bottom=med)
    ws.cell(row=r, column=4, value="=SUM(D6:D18)").number_format = "#,##0"
    ws.conditional_formatting.add("E6:E18", CellIsRule(operator="greaterThan", formula=["0"], fill=green, font=Font(color="006100")))
    ws.conditional_formatting.add("E6:E18", CellIsRule(operator="lessThan", formula=["0"], fill=red, font=Font(color="9C0006")))
    ws.conditional_formatting.add("J6:J18", FormulaRule(formula=['$J6="check"'], fill=red))       # relative ref (v2 ignored the offset)
    for col, w in zip("CDEFGHIJ", [16, 14, 9, 16, 9, 9, 12, 10]): ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("SLIDE_2 panel")                 # key/value control panel, no markers
    ws2["A1"] = "Liquidity panel"; ws2["A1"].font = Font(bold=True, size=14)
    rows = [("LCR", 1.523, "0.0%"), ("NSFR", 1.21, "0.0%"), ("Cash (EUR bn)", 12.345, "0.00"), ("Rating", "A+", None)]
    for k, (lab, v, f) in enumerate(rows):
        a = ws2.cell(row=3 + k * 2, column=2, value=lab); a.fill = navy; a.font = Font(color="FFFFFF", bold=True)
        b = ws2.cell(row=3 + k * 2, column=4, value=v); b.border = box
        if f: b.number_format = f
    ws2["B12"] = "Source: treasury"; ws2["B12"].font = Font(italic=True, size=9)
    ws2.column_dimensions["B"].width = 22; ws2.column_dimensions["D"].width = 14

    ws5 = wb.create_sheet("SLIDE_3 pair")                # two tables of different size side by side
    ws5["A1"] = "Loans vs deposits by region"
    for (r0, c0, nrow, ncol, w) in ((3, 2, 6, 3, 12), (3, 7, 9, 4, 9)):
        ws5.cell(row=r0, column=c0, value="x"); ws5.cell(row=r0 + nrow + 2, column=c0 + ncol + 1, value="x")
        for j in range(ncol):
            h = ws5.cell(row=r0 + 1, column=c0 + 1 + j, value=["Region", "Q1", "Q2", "Q3"][j]); h.font = Font(bold=True); h.fill = grey
            ws5.column_dimensions[ws5.cell(row=1, column=c0 + 1 + j).column_letter].width = w + 2 * j
        for i in range(nrow):
            ws5.cell(row=r0 + 2 + i, column=c0 + 1, value="R%d" % (i + 1))
            for j in range(1, ncol):
                ws5.cell(row=r0 + 2 + i, column=c0 + 1 + j, value=round((i - 3) * 37.5 + j * 11.25, 2)).number_format = "#,##0.0"
    ws3 = wb.create_sheet("Hidden data"); ws3["A1"] = "secret"; ws3.sheet_state = "hidden"
    ws4 = wb.create_sheet("Numbers")
    samples = [(1234.5, "#,##0.00"), (0.000123, "0.00E+00"), (1.5, "# ?/?"), (12345678901234, "General"), (0.1 + 0.2, "General"), (45000.75, "dd mmm yyyy hh:mm AM/PM"),
               (45000, "[$-410]mmmm yyyy"), (1234.5, "#,##0 [$€-410]"), (-5, "0;[Red]-0"), (0.25, "0%")]
    for k, (v, f) in enumerate(samples):
        c = ws4.cell(row=2 + k, column=2, value=v); c.number_format = f
    wb.save(os.path.join(HERE, "report.xlsx"))

def weekly():
    """two tables like the weekly banking report (made-up numbers): weeks, budget, Δ groups with Abs./% columns"""
    wb = Workbook(); ws = wb.active; ws.title = "LOANS_DEPOSITS"
    ws["A1"] = "Total Banks Loans & Deposits"; ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "(mln Euro at last fixed exchange rate)"
    banks = ["VUB", "PBZ", "BIB", "Alex", "CIB", "ISP SLO", "ISP RO", "ISP ALB", "ISP BiH", "EximBank", "Pravex"]
    blue = PatternFill("solid", fgColor="1F4E79"); wf = Font(bold=True, color="FFFFFF", italic=True)
    groups = ["Δ vs. Budget", "Δ vs. Prev. Week", "Δ vs. EoM Aug 2026", "Δ vs. Q2", "Δ vs. BoY"]
    def table(top, name, seed):
        ws.cell(row=top, column=2, value="x"); ws.cell(row=top + 17, column=18, value="x")
        r1, r2 = top + 1, top + 2
        for i, (h, d) in enumerate([("Week37", "11/09/26"), ("Week38", "18/09/26"), ("Week39", "25/09/26")]):
            for r, v in ((r1, h), (r2, d)):
                c = ws.cell(row=r, column=4 + i, value=v); c.fill = blue; c.font = wf; c.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=r1, start_column=7, end_row=r2, end_column=7)
        c = ws.cell(row=r1, column=7, value="Eom Budget"); c.fill = blue; c.font = wf; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for g, name_g in enumerate(groups):
            col = 8 + 2 * g
            ws.merge_cells(start_row=r1, start_column=col, end_row=r1, end_column=col + 1)
            c = ws.cell(row=r1, column=col, value=name_g); c.fill = blue; c.font = wf; c.alignment = Alignment(horizontal="center")
            for k, sub in enumerate(["Abs.", "%"]):
                c = ws.cell(row=r2, column=col + k, value=sub); c.fill = blue; c.font = wf; c.alignment = Alignment(horizontal="center")
        rows = []
        for k, b in enumerate(banks):
            base = 1000 * (k + 2) + seed * 37
            wk = [base - 30 * ((k * 7 + seed) % 5), base - 12 * ((k * 3 + seed) % 4), base]
            budget = None if b in ("EximBank", "Pravex") else base - ((k * 53 + seed * 11) % 190) + 80
            ref = [budget, wk[1], base - 140 + (k * 29 + seed) % 260, base - 220 + (k * 41) % 400, base - 600 + (k * 61) % 900]
            rows.append([b] + wk + [budget] + [x for v in ref for x in ((None, None) if v is None else (base - v, (base - v) / v))])
        tot = [name] + [sum(r[i] for r in rows if r[i] is not None) for i in range(1, 5)]
        for g in range(5):
            refsum = sum(r[3] - r[5 + 2 * g] for r in rows if r[5 + 2 * g] is not None)
            ab = sum(r[5 + 2 * g] for r in rows if r[5 + 2 * g] is not None); tot += [ab, ab / refsum]
        for k, vals in enumerate([tot] + rows):
            r = top + 4 + k + (1 if k else 0)
            for i, v in enumerate(vals):
                c = ws.cell(row=r, column=3 + i, value=v)
                if i == 0: c.font = Font(bold=k == 0)
                elif i >= 5: c.number_format = "+#,##0;-#,##0;0" if i % 2 == 1 else "+0.0%;-0.0%;0.0%"
                else: c.number_format = "#,##0"
                if k == 0: c.font = Font(bold=True)
    table(4, "TOTAL BANKS LOANS", 1)
    table(24, "TOTAL BANKS DEPOSITS", 2)
    for col, w in zip("ABC", [2, 3, 22]): ws.column_dimensions[col].width = w
    wb.save(os.path.join(HERE, "weekly.xlsx"))

def plain():
    wb = Workbook(); ws = wb.active; ws.title = "Data"
    ws["A1"] = "Products"
    for r in range(3, 12):
        for c in range(1, 6):
            ws.cell(row=r, column=c, value=(f"P{r}" if c == 1 else r * c * 10.5) if r > 3 else f"Col {c}")
    wb.save(os.path.join(HERE, "plain.xlsx"))

def big():
    """a sheet larger than the wizard preview (400 rows / 80 columns): the preview shows a note above
    the grid - selections and suggestions must still sit exactly on their cells"""
    wb = Workbook(); ws = wb.active; ws.title = "WM_ALL_STOCK_new"
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].hidden = True; ws.column_dimensions["C"].width = 30
    hdr = Font(bold=True, color="FFFFFF")
    ws["C5"] = "(mln Euro at last fixed exchange rate)"
    for j, t in enumerate(["Stock", "Stock", "Stock", "EOM Bdg"]):
        c = ws.cell(row=4, column=4 + j, value=t); c.fill = navy; c.font = hdr
        c = ws.cell(row=5, column=4 + j, value=f"0{j + 1}/09/26"); c.fill = navy; c.font = hdr
    ws.merge_cells("D3:G3"); ws["D3"] = "Total"
    for r in range(6, 450):
        ws.cell(row=r, column=3, value=f"Row {r}")
        for c in range(4, 8):
            ws.cell(row=r, column=c, value=r * c)
    ws.cell(row=2, column=90, value="far right")
    wb.save(os.path.join(HERE, "big.xlsx"))

import sys
for name in (sys.argv[1:] or ["report", "plain", "big", "weekly"]):
    globals()[name]()
print("ok")
