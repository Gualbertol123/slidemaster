"""Generates B54.xlsx (python3 make_b54.py; needs openpyxl, dev machines only): the owner's two sheets of B54 (2026-10-23),
cell for cell - "KPI" (image 16: four products, their current value, three reference levels and the three deltas as
Excel percentages, the notes) and "Volumes" (image 17: a long table of two products' weekly volumes and their EoM
budget). Used by the core tests of B54 (kpi-cards, chart-card) and the e2e spec B54-01."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
navy = PatternFill("solid", fgColor="13294B")
white = Font(bold=True, color="FFFFFF")


def header(ws, names):
    for i, h in enumerate(names):
        c = ws.cell(row=1, column=1 + i, value=h)
        c.fill = navy
        c.font = white


def kpi():
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"
    header(ws, ["Product", "Current", "Budget_reference", "Previous_week", "EoM_reference", "Delta_vs_Budget", "Delta_vs_prev_Week", "Delta_vs_EoM"])
    rows = [("Loans", 10, 9.900990099, 10.20408163, 9.708737864),
            ("Deposits", -10, -10.1010101, -9.803921569, -10.30927835),
            ("WM", 10, 9.900990099, 10.20408163, 9.708737864),
            ("Certificates", -10, -10.1010101, -9.803921569, -10.30927835)]
    for k, (p, cur, bdg, prev, eom) in enumerate(rows):
        r = 2 + k
        for i, v in enumerate([p, cur, bdg, prev, eom]):
            ws.cell(row=r, column=1 + i, value=v)
        for i, v in enumerate([0.01, -0.02, 0.03]):
            ws.cell(row=r, column=6 + i, value=v).number_format = "0.0%"
    ws["J1"] = "NOTES"
    for k, t in enumerate(["All four cards illustrative", "Replace current AND references", "Deltas calculate from levels", "Negative values are allowed"]):
        ws.cell(row=2 + k, column=10, value=t)
    for col, w in zip("ABCDEFGHIJ", [14, 10, 18, 16, 16, 18, 20, 16, 4, 30]):
        ws.column_dimensions[col].width = w
    vs = wb.create_sheet("Volumes")
    header(vs, ["Product", "Period", "Volume_EUR_m", "EoM_Budget_EUR_m"])
    data = [("Loans", "Q3", 11576, 11790), ("Loans", "W37", 11411, 11790), ("Loans", "W38", 13000, 11790), ("Loans", "W39", 12000, 11790),
            ("Deposits", "Q3", 8420, 8700), ("Deposits", "W37", 8300, 8700), ("Deposits", "W38", 9000, 8700), ("Deposits", "W39", 7000, 8700)]
    for k, row in enumerate(data):
        for i, v in enumerate(row):
            vs.cell(row=2 + k, column=1 + i, value=v)
    for col, w in zip("ABCD", [12, 10, 16, 20]):
        vs.column_dimensions[col].width = w
    wb.save(os.path.join(HERE, "B54.xlsx"))


if __name__ == "__main__":
    kpi()
