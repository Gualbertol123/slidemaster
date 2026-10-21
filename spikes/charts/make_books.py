"""Workbooks of the 09 Phase 0 pictures (dev tool, not shipped; run from the repository root):

    python spikes/charts/make_books.py

* spikes/charts/books/loans-weekly.xlsx – the owner's reference chart's numbers (09 §3.3): sheet Weekly (Q3, Week37 …
  Week39 and the loans volume), sheet Targets (Eom Budget, EoM Ambition and their targets), sheet Gallery (the series of
  the chart-type pictures). Plain values, no formulas (openpyxl writes no cached values for formulas).
* spikes/charts/books/bridge.xlsx – the scatter and waterfall charts' numbers (S14.4; copied to
  tests/corpus/workbooks/charts/): sheet Bridge (a budget bridge from 900 to 1.030 and a profit walk that crosses zero,
  each a column of steps with its start and end as totals), sheet Scatter (twelve months' volume against margin for
  Retail and Corporate).
* spikes/charts/books/report-hidden.xlsx – tests/corpus/workbooks/report.xlsx with rows 10–12 of SLIDE_1 hidden: F11 and
  F14 close rows up through the very path that closes Excel's own hidden rows (buildLayout), so this file shows exactly
  what a slide with those rows hidden will show.
* spikes/charts/books/report-drop.xlsx – the same with rows 10–12 and columns I:J (Since, Note) hidden: a delete-version.
The two report copies are made by editing the sheet XML in the zip (every other byte kept: styles, cached formula values)."""
import os
import re
import zipfile

import openpyxl
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
BOOKS = os.path.join(HERE, "books")
REPORT = os.path.join(HERE, "..", "..", "tests", "corpus", "workbooks", "report.xlsx")
os.makedirs(BOOKS, exist_ok=True)


def loans():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly"
    ws["B2"] = "Loans volume, EUR m"
    for col, (cat, v) in zip("CDEF", [("Q3", 11576), ("Week37", 11258), ("Week38", 11258), ("Week39", 11324)]):
        ws[col + "4"] = cat
        ws[col + "5"] = v
        ws[col + "5"].number_format = "#,##0"
    ws["B5"] = "Loans"
    t = wb.create_sheet("Targets")
    for col, (name, target) in zip("BC", [("Eom Budget", 11624), ("EoM Ambition", 12000)]):
        t[col + "2"] = name
        t[col + "3"] = target
        t[col + "3"].number_format = "#,##0"
    t["A3"] = "Target"
    g = wb.create_sheet("Gallery")
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
    rows = {"Retail": [420, 455, 470, 510, 498, 530], "Corporate": [310, 300, 335, 360, 372, 390],
            "SME": [150, 162, 158, 171, 180, 176], "Margin %": [2.1, 2.3, 2.2, 2.6, 2.5, 2.8]}
    for i, m in enumerate(months):
        g.cell(row=1, column=2 + i, value=m)
    for r, (name, vals) in enumerate(rows.items(), start=2):
        g.cell(row=r, column=1, value=name)
        for i, v in enumerate(vals):
            g.cell(row=r, column=2 + i, value=v)
    for c in g[1]:
        c.font = Font(bold=True)
    wb.save(os.path.join(BOOKS, "loans-weekly.xlsx"))


def bridge():
    wb = openpyxl.Workbook()
    b = wb.active
    b.title = "Bridge"
    b["B2"] = "Budget bridge, EUR m"
    b["C3"], b["E3"] = "Bridge", "Profit walk"
    rows = [("Start", 900, "Start", 120), ("Retail", 260, "Volume", 60), ("Corporate", 180, "Price", -150),
            ("Costs", -220, "Costs", -90), ("Tax", -90, "FX", 40), ("End", 1030, "End", -20)]
    for r, (c1, v1, c2, v2) in enumerate(rows, start=4):
        b["B%d" % r], b["C%d" % r], b["D%d" % r], b["E%d" % r] = c1, v1, c2, v2
        b["C%d" % r].number_format = b["E%d" % r].number_format = "#,##0;-#,##0"
    s = wb.create_sheet("Scatter")
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    vol = [730, 755, 805, 870, 870, 920, 905, 860, 940, 985, 1010, 1060]
    mar = [2.1, 2.3, 2.2, 2.6, 2.5, 2.8, 2.7, 2.4, 2.9, 3.0, 2.9, 3.2]
    cvol = [410, 430, 455, 470, 500, 520, 515, 540, 560, 585, 600, 640]
    cmar = [1.6, 1.7, 1.65, 1.8, 1.9, 1.95, 1.9, 2.05, 2.1, 2.2, 2.15, 2.3]
    for c, h in zip("ABCDE", ["Month", "Retail volume", "Retail margin %", "Corporate volume", "Corporate margin %"]):
        s[c + "1"] = h
        s[c + "1"].font = Font(bold=True)
    for r, row in enumerate(zip(months, vol, mar, cvol, cmar), start=2):
        for c, v in zip("ABCDE", row):
            s["%s%d" % (c, r)] = v
        s["B%d" % r].number_format = s["D%d" % r].number_format = "#,##0"
        s["C%d" % r].number_format = s["E%d" % r].number_format = "0.0"
    wb.save(os.path.join(BOOKS, "bridge.xlsx"))


def sheet_part(z):
    """the part of sheet SLIDE_1 (workbook.xml order → rels)"""
    wbx = z.read("xl/workbook.xml").decode("utf8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf8")
    rid = re.search(r'<sheet [^>]*name="SLIDE_1"[^>]*r:id="([^"]+)"', wbx).group(1)
    target = re.search(r'<Relationship [^>]*Id="%s"[^>]*Target="([^"]+)"' % rid, rels) or re.search(r'<Relationship [^>]*Target="([^"]+)"[^>]*Id="%s"' % rid, rels)
    t = target.group(1).lstrip("/")
    return t if t.startswith("xl/") else "xl/" + t


def hide(dst, rows, cols):
    with zipfile.ZipFile(REPORT) as z:
        part = sheet_part(z)
        xml = z.read(part).decode("utf8")
        for r in rows:
            xml, n = re.subn(r'<row r="%d"' % r, '<row r="%d" hidden="1"' % r, xml, count=1)
            assert n == 1, r
        for c in range(cols[0], cols[1] + 1) if cols else []:
            # report.xlsx defines each of these columns once (<col … min="9" max="9"/>): mark that element hidden
            xml, n = re.subn(r'<col (?![^>]*hidden)([^>]*min="%d" max="%d")' % (c, c), r'<col hidden="1" \1', xml, count=1)
            assert n == 1, c
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as out:
            for info in z.infolist():
                data = xml.encode("utf8") if info.filename == part else z.read(info.filename)
                out.writestr(info, data)


loans()
bridge()
hide(os.path.join(BOOKS, "report-hidden.xlsx"), [10, 11, 12], None)
hide(os.path.join(BOOKS, "report-drop.xlsx"), [10, 11, 12], (9, 10))
print("written:", sorted(os.listdir(BOOKS)))
