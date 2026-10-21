"""tools/pptxcheck.py (a dev tool, like pdfcheck.py): the real-objects gate (5) and the chart-formula check of (1)
(S14.P review, 1 and 10). The gate's counting is plain Python; the formula check reads a .pptx made with python-pptx,
skipped where python-pptx is not installed (the helper itself needs only the standard library)."""
import importlib.util
import os
import shutil
import tempfile
import unittest
import zipfile

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
spec = importlib.util.spec_from_file_location("sb_pptxcheck", os.path.join(REPO, "tools", "pptxcheck.py"))
pptxcheck = importlib.util.module_from_spec(spec)
spec.loader.exec_module(pptxcheck)
try:
    import pptx  # noqa: F401
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False


def slide(tables=0, charts=0, pictures=0):
    """a slide of content() (2) as the gate reads it"""
    return {"tables": tables, "charts": [{"type": "PIE"}] * charts, "pictures": pictures, "texts": []}


class RealObjectsTests(unittest.TestCase):
    def test_a_slide_whose_table_is_missing(self):
        # tests/corpus/decks/missing-table.json: slide 1 lists a table of Data and one of a deleted sheet, slide 2 only
        # the deleted sheet's. The core draws one table on slide 1 and none on slide 2, so the file has exactly that
        # (tools/real-objects.mjs lists only slide 1) and the gate passes
        content = {"slides": [slide(tables=1), slide()]}
        self.assertEqual(pptxcheck.real_objects(content, [{"slide": 1, "tables": 1, "charts": 0}], "python-pptx reads"), [])
        # the deck's JSON alone (a table per listed slide, S14.P's first gate) would have failed it
        self.assertEqual(pptxcheck.real_objects(content, [{"slide": 1, "tables": 1, "charts": 0}, {"slide": 2, "tables": 1, "charts": 0}], "python-pptx reads"),
                         ["slide 2: python-pptx reads 0 real table(s), the export draws 1"])

    def test_a_picture_in_place_of_a_table_or_chart_fails(self):
        content = {"slides": [slide(pictures=2)]}
        self.assertEqual(pptxcheck.real_objects(content, [{"slide": 1, "tables": 1, "charts": 1}], "python-pptx reads"),
                         ["slide 1: python-pptx reads 0 real table(s), the export draws 1", "slide 1: python-pptx reads 0 native chart(s), the export draws 1"])

    def test_exact_counts_and_slides_not_listed(self):
        content = {"slides": [slide(tables=2, charts=1), slide(tables=1)]}
        # two tables where the export draws one is wrong too; a slide the expectation does not list must hold none
        self.assertEqual(pptxcheck.real_objects(content, [{"slide": 1, "tables": 1, "charts": 1}], "python-pptx reads"),
                         ["slide 1: python-pptx reads 2 real table(s), the export draws 1", "slide 2: python-pptx reads 1 real table(s), the export draws 0"])
        self.assertEqual(pptxcheck.real_objects(content, [{"slide": 1, "tables": 2, "charts": 1}, {"slide": 2, "tables": 1}], "python-pptx reads"), [])
        # a slide expected beyond the file's last one (a hidden slide counted by mistake)
        self.assertEqual(pptxcheck.real_objects({"slides": []}, [{"slide": 1, "tables": 1, "charts": 0}], "python-pptx reads"),
                         ["slide 1: python-pptx reads 0 real table(s), the export draws 1"])
        # no content (python-pptx could not read the file): every expected object is missing; an empty expectation passes
        self.assertEqual(pptxcheck.real_objects(None, [], "python-pptx reads"), [])


@unittest.skipUnless(HAVE_PPTX, "python-pptx is not installed (pip install python-pptx==1.0.2)")
class ChartFormulaTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="sb-pptxcheck-")
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def deck(self, formula):
        """a one-slide deck with a native chart whose category formula is `formula`"""
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Emu
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ["a", "b"]
        data.add_series("S", (1, 2))
        s.shapes.add_chart(XL_CHART_TYPE.PIE, 0, 0, Emu(4000000), Emu(3000000), data)
        src = os.path.join(self.tmp, "src.pptx")
        prs.save(src)
        out = os.path.join(self.tmp, "deck-%d.pptx" % len(os.listdir(self.tmp)))
        with zipfile.ZipFile(src) as a, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as b:
            for i in a.infolist():
                x = a.read(i.filename)
                if i.filename.startswith("ppt/charts/chart") and i.filename.endswith(".xml"):
                    x = x.replace(b"<c:f>Sheet1!$A$2:$A$3</c:f>", ("<c:f>%s</c:f>" % formula).encode("utf-8"))
                    self.assertIn(("<c:f>%s</c:f>" % formula).encode("utf-8"), x)
                b.writestr(i, x)
        return out

    def file_problems(self, formula):
        bad, _ = pptxcheck.structure(self.deck(formula))
        return [p for p in bad if "names a file" in p]

    def test_own_workbook_and_a_sheet_named_like_a_file_pass(self):
        self.assertEqual(self.file_problems("Sheet1!$A$2:$A$3"), [])
        self.assertEqual(self.file_problems("'Q3.xlsx'!$A$2:$A$3"), [])
        self.assertEqual(self.file_problems("'Budget.xlsm results'!$A$2:$A$3"), [])

    def test_another_workbook_fails(self):
        for f in ("[1]Sheet1!$A$2:$A$3", "[Peers.xlsx]Peer!$A$2:$A$3", "'C:\\Data\\[Peers.xlsx]Peer'!$A$2:$A$3", "'/data/Peers.xlsx'#Peer!$A$2"):
            self.assertEqual(len(self.file_problems(f)), 1, f)


if __name__ == "__main__":
    unittest.main()
